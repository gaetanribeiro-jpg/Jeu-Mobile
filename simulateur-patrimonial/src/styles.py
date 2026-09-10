# -*- coding: utf-8 -*-
"""Styles, formats et helpers communs à toutes les feuilles du classeur."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment

POLICE = "Arial"

# --- Palette -----------------------------------------------------------------
NAVY = "1F3864"      # bandeaux de titre
BLEU = "2E5C9A"      # sous-sections
LIGHT = "D9E2F3"     # fond des sous-sections
CREAM = "FFF2CC"     # fond des cellules de saisie
GREY = "F2F2F2"      # fond des lignes de résultat
GREY2 = "E7E6E6"     # fond des totaux
WHITE = "FFFFFF"
ALERT = "FCE4D6"     # fond des alertes

C_INPUT = "0000FF"   # bleu : saisie en dur
C_CALC = "000000"    # noir : formule interne à la feuille
C_LINK = "008000"    # vert : formule pointant une autre feuille
C_WARN = "C00000"    # rouge : alerte
C_NOTE = "808080"    # gris : commentaire

# --- Formats de nombre -------------------------------------------------------
F_EUR = '#,##0 "€";[Red]-#,##0 "€";"-"'
F_EUR2 = '#,##0.00 "€";[Red]-#,##0.00 "€";"-"'
F_EUR4 = '#,##0.0000 "€"'
F_PCT = '0.0%;[Red]-0.0%;"-"'
F_PCT2 = '0.00%;[Red]-0.00%;"-"'
F_INT = '#,##0;[Red]-#,##0;"-"'
F_DEC = '#,##0.00'
F_DATE = 'DD/MM/YYYY'
F_TXT = '@'

_thin = Side(style="thin", color="BFBFBF")
_med = Side(style="medium", color=NAVY)
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
TOPLINE = Border(top=Side(style="thin", color=NAVY))


def A(col, row):
    """Adresse absolue $C$12 à partir d'une lettre (ou d'un index) et d'une ligne."""
    if isinstance(col, int):
        col = get_column_letter(col)
    return "${}${}".format(col, row)


def dn(wb, name, sheet, ref):
    """Enregistre un nom défini (plage nommée) dans le classeur."""
    attr = "{}!{}".format(sheet, ref)
    try:
        wb.defined_names.add(DefinedName(name, attr_text=attr))
    except AttributeError:  # openpyxl >= 3.1
        wb.defined_names[name] = DefinedName(name, attr_text=attr)


# --- Écriture de cellules ----------------------------------------------------
def _set(ws, addr, value, font, fill=None, fmt=None, align=None, border=None,
         wrap=False, comment=None):
    c = ws[addr]
    c.value = value
    c.font = font
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=align or "general",
                            vertical="center", wrap_text=wrap)
    if border:
        c.border = border
    if comment:
        cm = Comment(comment, "Simulateur patrimonial")
        cm.width, cm.height = 320, 130
        c.comment = cm
    return c


def titre(ws, row, texte, ncols=12, sous_titre=None):
    """Bandeau de titre en haut de feuille."""
    _set(ws, A(1, row).replace("$", ""), texte,
         Font(name=POLICE, size=15, bold=True, color=WHITE), fill=NAVY, align="left")
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[row].height = 30
    if sous_titre:
        _set(ws, "A{}".format(row + 1), sous_titre,
             Font(name=POLICE, size=9, italic=True, color=C_NOTE))
        ws.row_dimensions[row + 1].height = 14


def section(ws, row, texte, ncols=12):
    """Bandeau de section."""
    for col in range(1, ncols + 1):
        c = ws.cell(row=row, column=col)
        c.fill = PatternFill("solid", fgColor=BLEU)
        c.font = Font(name=POLICE, size=10, bold=True, color=WHITE)
    ws.cell(row=row, column=1).value = texte
    ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20


def entetes(ws, row, col_start, valeurs, largeurs=None, wrap=True):
    """Ligne d'en-têtes de tableau."""
    for i, v in enumerate(valeurs):
        col = col_start + i
        _set(ws, "{}{}".format(get_column_letter(col), row), v,
             Font(name=POLICE, size=9, bold=True, color=NAVY),
             fill=LIGHT, align="center", border=BOX, wrap=wrap)
        if largeurs:
            ws.column_dimensions[get_column_letter(col)].width = largeurs[i]
    ws.row_dimensions[row].height = 28


def lab(ws, addr, texte, bold=False, italic=False, size=10, color=None, indent=0):
    return _set(ws, addr, texte,
                Font(name=POLICE, size=size, bold=bold, italic=italic,
                     color=color or "000000"))


def inp(ws, addr, valeur=None, fmt=None, align=None, wrap=False, comment=None):
    """Cellule de saisie : police bleue sur fond crème."""
    return _set(ws, addr, valeur, Font(name=POLICE, size=10, color=C_INPUT),
                fill=CREAM, fmt=fmt, align=align, border=BOX, wrap=wrap,
                comment=comment)


def calc(ws, addr, formule, fmt=None, bold=False, align=None, comment=None):
    """Cellule calculée à l'intérieur de la feuille : police noire."""
    return _set(ws, addr, formule, Font(name=POLICE, size=10, bold=bold, color=C_CALC),
                fmt=fmt, align=align, comment=comment)


def link(ws, addr, formule, fmt=None, bold=False, align=None, comment=None):
    """Cellule calculée à partir d'une autre feuille : police verte."""
    return _set(ws, addr, formule, Font(name=POLICE, size=10, bold=bold, color=C_LINK),
                fmt=fmt, align=align, comment=comment)


def total(ws, addr, formule, fmt=None, align=None):
    """Ligne de total : gras, fond gris."""
    return _set(ws, addr, formule, Font(name=POLICE, size=10, bold=True, color=C_CALC),
                fill=GREY2, fmt=fmt, align=align, border=BOX)


def resultat(ws, addr, formule, fmt=None, align=None):
    """Résultat mis en avant : gras, blanc sur fond marine."""
    return _set(ws, addr, formule, Font(name=POLICE, size=11, bold=True, color=WHITE),
                fill=NAVY, fmt=fmt, align=align or "center", border=BOX)


def note(ws, addr, texte, ncols=1):
    c = _set(ws, addr, texte, Font(name=POLICE, size=8, italic=True, color=C_NOTE),
             wrap=False)
    return c


def alerte(ws, addr, formule):
    return _set(ws, addr, formule, Font(name=POLICE, size=10, bold=True, color=C_WARN),
                fill=ALERT, align="left", border=BOX, wrap=True)


def alerte_large(ws, addr, formule, span=4, hauteur=44):
    """Alerte occupant plusieurs colonnes (texte long)."""
    col = "".join(ch for ch in addr if ch.isalpha())
    row = int("".join(ch for ch in addr if ch.isdigit()))
    fin = get_column_letter(ws[addr].column + span - 1)
    c = alerte(ws, addr, formule)
    ws.merge_cells("{}{}:{}{}".format(col, row, fin, row))
    ws.row_dimensions[row].height = hauteur
    return c


def dv_liste(ws, plage, options=None, source=None, titre_msg=None, msg=None):
    """Ajoute une validation par liste sur une plage ('C4:C10')."""
    if source:
        f1 = source
    else:
        f1 = '"{}"'.format(",".join(options))
    d = DataValidation(type="list", formula1=f1, allow_blank=True, showDropDown=False)
    d.errorTitle = "Valeur non autorisée"
    d.error = "Choisissez une valeur dans la liste déroulante."
    if titre_msg:
        d.promptTitle, d.prompt, d.showInputMessage = titre_msg, msg or "", True
    ws.add_data_validation(d)
    d.add(plage)
    return d


def largeurs(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def mise_en_page(ws, zoom=100, freeze=None, paysage=True, ncols=12):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = zoom
    if freeze:
        ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape" if paysage else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.tabColor = NAVY
